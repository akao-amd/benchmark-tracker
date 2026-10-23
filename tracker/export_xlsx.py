"""Write the workbook Excel is asked to read.

The sheet matches the circulated comparison layout: stacked blocks, summary
formulas into a raw block at column M, ratios and GEOMEAN when a block names
a baseline, and one scatter per shape for blocks marked for the figure.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from .archive import Archive, Block
from .bench import RAW_HEADERS

SUMMARY_COL = 2
RAW_COL = 13
SUMMARY = (
    ("input_len", "", None),
    ("output_len", "", None),
    ("concurrency", "", None),
    ("Throughput per GPU", "(tok/s/GPU)", "0"),
    ("Interactivity", "(tok/s/user)", "0.00"),
    ("Median TTFT", "(ms)", "0.0"),
    ("Median TPOT", "(ms)", "0.0"),
)
RATIO_HEADERS = ("tok/s/gpu %", "TTFT %", "TPOT %")


class ExportError(Exception):
    pass


@dataclass
class ShapeLayout:
    shape: tuple[int, int]
    rows: list[tuple[int, int]]
    geomean_row: int | None


@dataclass
class BlockLayout:
    block: Block
    top: int
    next_top: int
    shapes: list[ShapeLayout] = field(default_factory=list)

    def row_of(self, input_len: int, output_len: int, concurrency: int) -> int | None:
        for shape in self.shapes:
            if shape.shape != (input_len, output_len):
                continue
            for conc, row in shape.rows:
                if conc == concurrency:
                    return row
        return None


def _plan(block: Block, top: int, has_ratios: bool) -> BlockLayout | None:
    points = block.included_points()
    if not points:
        return None
    grouped: list[tuple[tuple[int, int], list]] = []
    for point in points:
        if not grouped or grouped[-1][0] != point.shape:
            grouped.append((point.shape, [point]))
        else:
            grouped[-1][1].append(point)
    row = top + 5
    shapes = []
    for index, (shape, bucket) in enumerate(grouped):
        placed = []
        for point in bucket:
            placed.append((point.concurrency, row))
            row += 1
        geomean_row = None
        if has_ratios:
            geomean_row = row
            row += 1
        shapes.append(ShapeLayout(shape, placed, geomean_row))
        if index != len(grouped) - 1:
            row += 1
    return BlockLayout(block, top, row + 2, shapes)


def _tp_text(value: str) -> str:
    text = value.strip()
    if text.upper().startswith("TP"):
        return text
    if text.isdigit():
        return f"TP{text}"
    return text


def _raw(point, block: Block, model: str) -> dict:
    raw = {header: None for header in RAW_HEADERS}
    for key, value in point.extra.items():
        if key in raw:
            raw[key] = value
    raw["ISL"] = point.input_len
    raw["OSL"] = point.output_len
    raw["Concurrency"] = point.concurrency
    raw["Throughput/GPU (tok/s)"] = point.server_total / block.gpus
    if point.server_output is not None:
        raw["Output Throughput/GPU (tok/s)"] = point.server_output / block.gpus
    if point.server_input is not None:
        raw["Input Throughput/GPU (tok/s)"] = point.server_input / block.gpus
    raw["Median TTFT (ms)"] = point.median_ttft_ms
    raw["Median TPOT (ms)"] = point.median_tpot_ms
    raw["Median Interactivity (tok/s/user)"] = point.interactivity()
    if model:
        raw["Model"] = model
    if block.attr("chip"):
        raw["Hardware"] = block.attr("chip")
    if block.attr("framework"):
        raw["Framework"] = block.attr("framework")
    if block.attr("tp"):
        raw["TP"] = _tp_text(block.attr("tp"))
    return raw


def _sheet_title(title: str, taken: list[str]) -> str:
    """Excel sheet names are at most 31 characters and reject \\ / * ? : [ ]."""
    cleaned = " ".join((title or "Benchmark").translate(str.maketrans("\\/*?:[]", "       ")).split())
    cleaned = cleaned[:31] or "Benchmark"
    if cleaned not in taken:
        return cleaned
    number = 2
    while True:
        suffix = f" {number}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        if candidate not in taken:
            return candidate
        number += 1


def _place_second(book, sheet) -> None:
    """The new tab sits immediately to the right of the leftmost tab."""
    names = list(book.sheetnames)
    if len(names) < 2 or sheet.title not in names:
        return
    book.move_sheet(sheet, offset=1 - names.index(sheet.title))


def append_workbook(existing: bytes, archive: Archive, title: str) -> tuple[bytes, str]:
    """Add one benchmark sheet to an existing workbook. Other sheets stay."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ExportError("openpyxl is required to write an xlsx file") from exc
    try:
        book = load_workbook(BytesIO(existing))
    except Exception as exc:
        raise ExportError(f"the workbook could not be opened: {exc}") from exc
    before = list(book.sheetnames)
    data = workbook_bytes(archive, book=book, sheet_title=title)
    added = [name for name in book.sheetnames if name not in before]
    missing = [name for name in before if name not in book.sheetnames]
    if missing or len(added) != 1:
        raise ExportError("the workbook would have lost an existing sheet, so it was not changed")
    return data, added[0]


def workbook_bytes(archive: Archive, book=None, sheet_title: str = "Benchmark") -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import Reference, ScatterChart, Series
        from openpyxl.chart.axis import ChartLines
        from openpyxl.chart.layout import Layout, ManualLayout
        from openpyxl.chart.marker import Marker
        from openpyxl.chart.series import SeriesLabel
        from openpyxl.chart.shapes import GraphicalProperties
        from openpyxl.chart.text import RichText
        from openpyxl.drawing.line import LineProperties
        from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ExportError("openpyxl is required to write an xlsx file") from exc

    thin = Border(
        left=Side(style="thin", color="D6D3D1"),
        right=Side(style="thin", color="D6D3D1"),
        top=Side(style="thin", color="D6D3D1"),
        bottom=Side(style="thin", color="D6D3D1"),
    )
    font = Font(name="Arial", size=11)
    bold = Font(name="Arial", size=11, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    green = PatternFill("solid", fgColor="92D050")
    ratio_fill = PatternFill("solid", fgColor="FEF3C7")
    geo_fill = PatternFill("solid", fgColor="F5F5F4")
    helper_fill = PatternFill("solid", fgColor="E0F2FE")

    def put(cell, value, *, face=font, fill=None, align=center, fmt=None):
        cell.value = "" if value is None else value
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"
        if fill is not None:
            cell.fill = fill
        cell.font = face
        cell.alignment = align
        cell.border = thin
        if fmt:
            cell.number_format = fmt

    def formula(cell, text, **kwargs):
        put(cell, text, **kwargs)
        cell.data_type = "f"

    if book is None:
        book = Workbook()
        sheet = book.active
        sheet.title = _sheet_title(sheet_title, [])
    else:
        sheet = book.create_sheet(_sheet_title(sheet_title, book.sheetnames))
        _place_second(book, sheet)
    book.active = sheet
    by_id = {block.id: block for block in archive.blocks}

    def baseline_of(block: Block) -> Block | None:
        if block.baseline or not block.baseline_id:
            return None
        target = by_id.get(block.baseline_id)
        if target is None or not target.baseline:
            return None
        return target

    layouts = []
    cursor = 2
    for block in archive.blocks:
        base = baseline_of(block)
        layout = _plan(block, cursor, base is not None and bool(base.included_points()))
        if layout is None:
            continue
        layouts.append(layout)
        cursor = layout.next_top
    layout_of = {layout.block.id: layout for layout in layouts}

    if not layouts:
        sheet["B2"] = "No measurements to export."
        sheet.sheet_view.showGridLines = False
        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    points_of = {block.id: { (point.input_len, point.output_len, point.concurrency): point for point in block.included_points() } for block in archive.blocks}

    for layout in layouts:
        block = layout.block
        baseline = baseline_of(block)
        has_ratios = baseline is not None and baseline.id in layout_of
        top = layout.top
        chip = block.attr("chip")
        framework = block.attr("framework")
        tp = _tp_text(block.attr("tp")) if block.attr("tp") else ""
        header_label = chip or block.label
        put(sheet.cell(top, SUMMARY_COL, header_label), header_label, face=bold, fill=green if header_label else None)
        put(sheet.cell(top, SUMMARY_COL + 1, framework), framework, face=bold, fill=green if framework else None)
        put(sheet.cell(top, SUMMARY_COL + 2, tp), tp, face=bold)
        if has_ratios:
            versus = f"{block.label} v.s. {baseline.label}"
            put(sheet.cell(top, SUMMARY_COL + 7, versus), versus, face=bold, fill=ratio_fill, align=left)
            sheet.merge_cells(start_row=top, start_column=SUMMARY_COL + 7, end_row=top, end_column=SUMMARY_COL + 9)
        docker = block.attr("docker image")
        put(sheet.cell(top + 1, SUMMARY_COL, "Docker" if docker else ""), "Docker" if docker else "", face=bold, align=left)
        put(sheet.cell(top + 1, SUMMARY_COL + 1, docker), docker, align=left)
        if docker:
            sheet.merge_cells(start_row=top + 1, start_column=SUMMARY_COL + 1, end_row=top + 1, end_column=SUMMARY_COL + 6)
        machine = block.attr("host name")
        accuracy = block.attr("accuracy score")
        put(sheet.cell(top + 2, SUMMARY_COL, "Machine" if machine else ""), "Machine" if machine else "", face=bold, align=left)
        put(sheet.cell(top + 2, SUMMARY_COL + 1, machine), machine, align=left)
        accuracy_value: object = accuracy
        if accuracy:
            try:
                accuracy_value = float(accuracy)
            except ValueError:
                accuracy_value = accuracy
        put(sheet.cell(top + 2, SUMMARY_COL + 3, "Accuracy" if accuracy else ""), "Accuracy" if accuracy else "", face=bold, align=left)
        put(sheet.cell(top + 2, SUMMARY_COL + 4, accuracy_value if accuracy else ""), accuracy_value if accuracy else "", align=left)
        dates = [point.extra.get("Date") for point in block.included_points() if point.extra.get("Date")]
        if dates:
            put(sheet.cell(top + 2, SUMMARY_COL + 5, "Date"), "Date", face=bold, align=left)
            put(sheet.cell(top + 2, SUMMARY_COL + 6, dates[0]), dates[0], align=left)
        bkc = block.attr("BKC")
        if bkc:
            put(sheet.cell(top + 2, SUMMARY_COL + 7, "BKC"), "BKC", face=bold, align=left)
            put(sheet.cell(top + 2, SUMMARY_COL + 8, bkc), bkc, align=left)
        header_row = top + 3
        unit_row = top + 4
        sheet.row_dimensions[header_row].height = 32
        for offset, (header, unit, fmt) in enumerate(SUMMARY):
            put(sheet.cell(header_row, SUMMARY_COL + offset, header), header, face=bold)
            if unit:
                put(sheet.cell(unit_row, SUMMARY_COL + offset, unit), unit)
            letter = get_column_letter(SUMMARY_COL + offset)
            sheet.column_dimensions[letter].width = max(14, len(header) + 2)
        if has_ratios:
            for offset, header in enumerate(RATIO_HEADERS):
                put(sheet.cell(header_row, SUMMARY_COL + 7 + offset, header), header, face=bold, fill=ratio_fill)
                sheet.column_dimensions[get_column_letter(SUMMARY_COL + 7 + offset)].width = 14
        for offset, header in enumerate(RAW_HEADERS):
            put(sheet.cell(header_row, RAW_COL + offset, header), header, face=bold, fill=helper_fill)
            sheet.column_dimensions[get_column_letter(RAW_COL + offset)].width = max(12, min(28, len(header) + 1))
        raw_index = {header: RAW_COL + index for index, header in enumerate(RAW_HEADERS)}
        tput_col = get_column_letter(SUMMARY_COL + 3)
        ttft_col = get_column_letter(SUMMARY_COL + 5)
        tpot_col = get_column_letter(SUMMARY_COL + 6)
        base_layout = layout_of.get(baseline.id) if has_ratios else None
        for shape in layout.shapes:
            for conc, row in shape.rows:
                point = points_of[block.id][(shape.shape[0], shape.shape[1], conc)]
                raw = _raw(point, block, archive.model)
                for offset, header in enumerate(RAW_HEADERS):
                    value = raw.get(header)
                    fmt = "0.00" if isinstance(value, float) else None
                    put(sheet.cell(row, RAW_COL + offset), None if value in (None, "") else value, fill=helper_fill, fmt=fmt)
                values = [
                    point.input_len,
                    point.output_len,
                    conc,
                    f"={get_column_letter(raw_index['Throughput/GPU (tok/s)'])}{row}",
                    f'=IF({tpot_col}{row}="","",1000/{tpot_col}{row})',
                    f"={get_column_letter(raw_index['Median TTFT (ms)'])}{row}",
                    f"={get_column_letter(raw_index['Median TPOT (ms)'])}{row}",
                ]
                for offset, value in enumerate(values):
                    cell = sheet.cell(row, SUMMARY_COL + offset)
                    fmt = SUMMARY[offset][2]
                    if isinstance(value, str) and value.startswith("="):
                        formula(cell, value, fmt=fmt)
                    else:
                        put(cell, value, fmt=fmt)
                if has_ratios and base_layout is not None:
                    base_row = base_layout.row_of(shape.shape[0], shape.shape[1], conc)
                    if base_row is not None:
                        for offset, column in enumerate((tput_col, ttft_col, tpot_col)):
                            formula(
                                sheet.cell(row, SUMMARY_COL + 7 + offset),
                                f"={column}{row}/{column}{base_row}",
                                fill=ratio_fill,
                                fmt="0%",
                            )
            matched = [
                base_layout.row_of(shape.shape[0], shape.shape[1], conc) is not None
                for conc, _row in shape.rows
            ] if has_ratios and base_layout is not None else []
            if shape.geomean_row is not None and any(matched):
                first = shape.rows[0][1]
                last = shape.rows[-1][1]
                for offset in range(3):
                    column = get_column_letter(SUMMARY_COL + 7 + offset)
                    formula(
                        sheet.cell(shape.geomean_row, SUMMARY_COL + 7 + offset),
                        f"=GEOMEAN({column}{first}:{column}{last})",
                        face=bold,
                        fill=geo_fill,
                        fmt="0%",
                    )
                put(sheet.cell(shape.geomean_row, SUMMARY_COL + 2, "geomean"), "geomean", face=bold, fill=geo_fill)

    _charts(
        sheet,
        archive,
        layouts,
        cursor,
        ScatterChart,
        Series,
        SeriesLabel,
        Reference,
        ChartLines,
        Layout,
        ManualLayout,
        Marker,
        GraphicalProperties,
        RichText,
        LineProperties,
        CharacterProperties,
        Paragraph,
        ParagraphProperties,
        RichTextProperties,
        get_column_letter,
    )
    sheet.column_dimensions["A"].width = 3
    sheet.freeze_panes = None
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _charts(sheet, archive, layouts, content_bottom, ScatterChart, Series, SeriesLabel, Reference, ChartLines, Layout, ManualLayout, Marker, GraphicalProperties, RichText, LineProperties, CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties, get_column_letter):
    sheet_ink = "404040"
    title_ink = "D9D9D9"
    ink = "BFBFBF"
    grid = "8C8C8C"
    series_line = ("4F81BD", "C0504D", "9BBB59", "8064A2", "4BACC6", "F79646")
    series_marker = ("95B3D7", "D99694", "C3D69B", "B3A2C7", "93CDDD", "FAC090")

    def rich_text(size, color, bold, rot=None):
        char = CharacterProperties(sz=size * 100, b=bold, solidFill=color)
        return RichText(bodyPr=RichTextProperties(rot=rot, anchor="ctr"), p=[Paragraph(pPr=ParagraphProperties(defRPr=char))])

    def paint_title(title, size, color, bold, rot=None):
        title.tx.rich.bodyPr = RichTextProperties(rot=rot, anchor="ctr")
        title.tx.rich.paragraphs[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=size * 100, b=bold, solidFill=color))
        title.txPr = rich_text(size, color, bold, rot)

    def box(left, top, right, bottom, inner=False):
        return Layout(manualLayout=ManualLayout(
            layoutTarget="inner" if inner else None,
            xMode="edge", yMode="edge", wMode="edge", hMode="edge",
            x=left, y=top, w=right, h=bottom,
        ))

    def show_axis(axis, position, number_format):
        axis.delete = False
        axis.axPos = position
        axis.tickLblPos = "nextTo"
        axis.number_format = number_format
        axis.majorTickMark = "out"
        axis.txPr = rich_text(9, ink, False)
        axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=ink, w=6350))
        axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=grid, w=6350)))

    charted = [layout for layout in layouts if layout.block.chart]
    if not charted:
        return
    figures = archive.present().get("figures") or []
    bottom = content_bottom + 2
    for index, figure in enumerate(figures):
        shape = (figure["input_len"], figure["output_len"])
        chart = ScatterChart()
        chart.scatterStyle = "lineMarker"
        chart.x_axis.title = "Interactivity (token/s/user)"
        chart.y_axis.title = "Total Token Throughput per GPU (token/s/#gpu)"
        chart.title = figure["title"]
        show_axis(chart.x_axis, "b", "0")
        show_axis(chart.y_axis, "l", "0")
        chart.width = 18
        chart.height = 11
        drawn = 0
        for series_index, layout in enumerate(charted):
            rows = []
            for item in layout.shapes:
                if item.shape == shape:
                    rows = [row for _conc, row in item.rows]
            if not rows:
                continue
            x_values = Reference(sheet, min_col=SUMMARY_COL + 4, min_row=rows[0], max_row=rows[-1])
            y_values = Reference(sheet, min_col=SUMMARY_COL + 3, min_row=rows[0], max_row=rows[-1])
            series = Series(y_values, x_values, title_from_data=False)
            series.title = SeriesLabel(v=layout.block.label)
            line = series_line[series_index % len(series_line)]
            marker = series_marker[series_index % len(series_marker)]
            series.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=line, w=22225))
            series.marker = Marker(symbol="circle", size=3, spPr=GraphicalProperties(solidFill=marker, ln=LineProperties(solidFill=line, w=6350)))
            series.smooth = False
            chart.series.append(series)
            drawn += 1
        if drawn == 0:
            continue
        chart.style = 2
        chart.graphical_properties = GraphicalProperties(solidFill=sheet_ink)
        chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True)
        paint_title(chart.title, 14, title_ink, True)
        paint_title(chart.x_axis.title, 9, ink, True)
        paint_title(chart.y_axis.title, 9, ink, True, rot=-5400000)
        chart.legend.position = "t"
        chart.legend.overlay = False
        chart.legend.txPr = rich_text(9, ink, False)
        chart.legend.graphicalProperties = GraphicalProperties(noFill=True)
        chart.legend.layout = box(0.05, 0.085, 0.95, 0.205)
        chart.layout = box(0.115, 0.235, 0.975, 0.835, inner=True)
        anchor = f"{get_column_letter(2 + (index % 2) * 12)}{bottom + (index // 2) * 18}"
        sheet.add_chart(chart, anchor)
