"""Measurements, the .bta file, the sheet, and the HTTP actions."""

import io
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from tracker.archive import Archive, ArchiveError, geomean
from tracker.bench import parse_name
from tracker.excel_push import describe, place_request
from tracker.export_xlsx import append_workbook, workbook_bytes
from tracker.fsview import list_dir, normalize
from tracker.ingest import scan_path
from tracker.onedrive import (
    OneDriveError,
    WorkbookLocked,
    _http_message,
    _open_for_edit,
    encode_share_url,
    file_name,
    remember_url,
)
from tracker.server import App, serve
from tracker.workspace import Workspace

ROOT = Path(__file__).resolve().parents[1]
REAL = Path("/home/akao/work-journal/2026/ww43/1022_fixlen_g01/0001_fixlen_sweep/L70_tp1")
R1_JSON = Path("/home/akao/work-journal/2026/ww43/1023_m355_r1/0001_r1_final/r1_json")


def bench_row(**overrides):
    row = {
        "total_input_tokens": 800,
        "total_output_tokens": 200,
        "duration": 1.0,
        "median_ttft_ms": 10.0,
        "median_tpot_ms": 20.0,
        "max_concurrency": 4,
        "input_len": 1024,
        "output_len": 1024,
        "backend": "sglang",
        "model_id": "org/gpt-oss-120b",
        "date": "20261022-191243",
    }
    row.update(overrides)
    row["total_token_throughput"] = (row["total_input_tokens"] + row["total_output_tokens"]) / row["duration"]
    row["output_throughput"] = row["total_output_tokens"] / row["duration"]
    return row


def write_bench(directory, label, isl, osl, conc, run=1, **overrides):
    row = bench_row(input_len=isl, output_len=osl, max_concurrency=conc, **overrides)
    path = Path(directory) / f"{label}_isl{isl}_osl{osl}_c{conc}_r{run}.json"
    path.write_text(json.dumps(row), encoding="utf-8")
    return path


class ModelTest(unittest.TestCase):
    def test_geomean_matches_excel(self):
        self.assertEqual(geomean([2, 8]), 4)
        self.assertIsNone(geomean([]))

    def test_scan_groups_label_and_skips_bad_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "tp1", 1024, 1024, 4)
            write_bench(tmp, "tp1", 1024, 1024, 8)
            write_bench(tmp, "tp1", 8192, 1024, 4, run=2)
            write_bench(tmp, "tp2", 1024, 1024, 4, total_input_tokens=100, total_output_tokens=50, duration=2)
            bad = bench_row(input_len=1024, output_len=1024, max_concurrency=16)
            bad["total_token_throughput"] = 1
            (Path(tmp) / "tp1_isl1024_osl1024_c16_r1.json").write_text(json.dumps(bad), encoding="utf-8")
            (Path(tmp) / "tp1_isl1024_osl1024_c32_r1.json").write_text(
                json.dumps(bench_row(input_len=2048, output_len=1024, max_concurrency=32)), encoding="utf-8"
            )
            (Path(tmp) / "tp1_isl1024_osl1024_c4_r1_OFFSPEC.json").write_text("{}", encoding="utf-8")
            log = Path(tmp) / "tp1_isl1024_osl1024_c64_r1.log"
            log.write_text(
                "Total token throughput (tok/s): 500\nMedian TTFT (ms): 12.5\nMedian TPOT (ms): 40\n",
                encoding="utf-8",
            )
            (Path(tmp) / "tp1_isl8192_osl1024_c4_r1_OFFSPEC.json").write_text(
                json.dumps(bench_row(input_len=8192, output_len=1024, max_concurrency=4)), encoding="utf-8"
            )
            blocks, notices, model = scan_path(Path(tmp))
            by = {block.label: block for block in blocks}
            self.assertEqual(model, "gpt-oss-120b")
            self.assertEqual(sorted(by), ["tp1", "tp2"])
            self.assertEqual(by["tp1"].active_run, 1)
            self.assertEqual(sorted(point.concurrency for point in by["tp1"].points if point.run == 1), [4, 4, 8, 64])
            self.assertTrue(any(point.input_len == 8192 and point.run == 1 for point in by["tp1"].points))
            self.assertTrue(any("throughput check failed" in item["message"] for item in notices))
            self.assertTrue(any("disagrees" in item["message"] for item in notices))
            self.assertTrue(any("OFFSPEC" in item["path"] for item in notices))
            self.assertEqual(by["tp1"].gpus, 1)
            self.assertEqual(by["tp1"].attributes["framework"], "sglang")

    def test_merge_choice_picks_the_row(self):
        def merged(choice, source_row, target_row, source_gpus=1, target_gpus=1):
            with tempfile.TemporaryDirectory() as tmp:
                write_bench(tmp, "src", 1024, 1024, 4, **source_row)
                write_bench(tmp, "src", 8192, 1024, 8)
                write_bench(tmp, "dst", 1024, 1024, 4, **target_row)
                blocks, _notices, _model = scan_path(Path(tmp))
                by = {block.label: block for block in blocks}
                by["src"].gpus = source_gpus
                by["dst"].gpus = target_gpus
                archive = Archive(blocks=[by["dst"], by["src"]])
                archive.merge(by["src"].id, by["dst"].id, choice)
                self.assertEqual([block.label for block in archive.blocks], ["dst"])
                self.assertEqual(sum(point.input_len == 8192 for point in archive.blocks[0].points), 1)
                matched = [point for point in archive.blocks[0].points if point.input_len == 1024]
                self.assertEqual(len(matched), 1)
                return matched[0].server_total

        low = {"total_input_tokens": 100, "total_output_tokens": 100}
        high = {"total_input_tokens": 400, "total_output_tokens": 100}
        self.assertEqual(merged("keep", low, high), 500)
        self.assertEqual(merged("overwrite", low, high), 200)
        self.assertEqual(merged("better", low, high), 500)
        self.assertEqual(merged("worse", low, high), 200)
        self.assertEqual(merged("better", {"total_input_tokens": 300, "total_output_tokens": 100}, {"total_input_tokens": 200, "total_output_tokens": 100}, 2, 1), 300)
        self.assertEqual(merged("worse", {"total_input_tokens": 300, "total_output_tokens": 100}, {"total_input_tokens": 200, "total_output_tokens": 100}, 2, 1), 400)

    def test_rescan_drops_a_stale_skip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = write_bench(tmp, "tp1", 1024, 1024, 4)
            archive = Archive()
            archive.add_notice(str(path), "name does not match <label>_isl<isl>_osl<osl>_c<conc>_r<run><suffix>.json", "skip")
            blocks, notices, model = scan_path(Path(tmp))
            archive.apply_scan(blocks, notices, {}, model)
            self.assertEqual(archive.notices, [])
            self.assertEqual(archive.blocks[0].points[0].source, str(path))

    def test_log_accepts_either_token_capitalization(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "tp1_isl1024_osl1024_c16_r1.log"
            path.write_text(
                "Total Token throughput (tok/s):          2629.18   \n"
                "Median TTFT (ms):                        74.56     \n"
                "Median TPOT (ms):                        11.36     \n",
                encoding="utf-8",
            )
            blocks, notices, _model = scan_path(Path(tmp))
            self.assertEqual(len(blocks), 1)
            self.assertEqual(blocks[0].points[0].server_total, 2629.18)
            self.assertFalse(any("missing total token" in item["message"] for item in notices))

    def test_suffix_after_the_run_is_accepted(self):
        plain = parse_name("tp1_isl1024_osl1024_c16_r1.json")
        marked = parse_name("g6_isl70000_osl300_c64_r1_OFFSPEC.json")
        dotted = parse_name("g7_isl70000_osl300_c64_r2.extra.log")
        self.assertEqual(plain["suffix"], "")
        self.assertEqual(marked["label"], "g6")
        self.assertEqual(marked["run"], 1)
        self.assertEqual(marked["suffix"], "_OFFSPEC")
        self.assertEqual(dotted["suffix"], ".extra")
        omitted = parse_name("g6_isl70000_osl300_c64_OFFSPEC.json")
        self.assertEqual(omitted["run"], 1)
        self.assertEqual(omitted["suffix"], "_OFFSPEC")
        self.assertIsNone(parse_name("g6_isl70000_osl300.json"))
        doubled = parse_name("R1__TP1__isl70000_osl300__c16.json")
        self.assertEqual(doubled["label"], "R1__TP1")
        self.assertEqual((doubled["input_len"], doubled["output_len"], doubled["concurrency"]), (70000, 300, 16))
        self.assertEqual(doubled["run"], 1)
        self.assertEqual(doubled["suffix"], "")

    def test_ratios_runs_and_missing_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = write_bench(tmp, "base", 1024, 1024, 4, total_input_tokens=50, total_output_tokens=50)
            write_bench(tmp, "base", 1024, 1024, 8, total_input_tokens=50, total_output_tokens=50)
            write_bench(tmp, "cand", 1024, 1024, 4, total_input_tokens=150, total_output_tokens=50)
            write_bench(tmp, "cand", 1024, 1024, 8, total_input_tokens=750, total_output_tokens=50)
            write_bench(tmp, "cand", 1024, 1024, 4, run=2, total_input_tokens=10, total_output_tokens=10)
            blocks, _notices, _model = scan_path(Path(tmp))
            archive = Archive(model="gpt-oss-120b", blocks=blocks)
            base, cand = archive.blocks
            base.baseline = True
            cand.baseline_id = base.id
            cand.chart = True
            base.chart = True
            view = archive.present()
            shape = view["blocks"][1]["shapes"][0]
            ratios = [row["ratios"]["tput"] for row in shape["rows"]]
            self.assertEqual(ratios, [2, 8])
            self.assertAlmostEqual(shape["geomean"]["tput"], 4)
            self.assertEqual(view["blocks"][1]["other_runs"], [2])
            self.assertEqual(len(view["figures"]), 1)
            self.assertIn("gpt-oss-120b", view["figures"][0]["title"])
            path.unlink()
            again = archive.present()
            self.assertTrue(again["blocks"][0]["shapes"][0]["rows"][0]["missing"])
            self.assertGreater(again["blocks"][0]["shapes"][0]["rows"][0]["per_gpu"], 0)

    def test_roundtrip_and_collision(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "tp1", 1024, 1024, 4)
            blocks, notices, model = scan_path(Path(tmp))
            archive = Archive(blocks=blocks, notices=notices, model=model)
            archive.blocks[0].attributes["docker image"] = "rocm/sgl-dev:test"
            archive.blocks[0].attributes["=odd"] = "kept"
            loaded = Archive.load(archive.dumps())
            self.assertEqual(loaded.blocks[0].attributes["docker image"], "rocm/sgl-dev:test")
            self.assertEqual(loaded.blocks[0].points[0].server_total, blocks[0].points[0].server_total)
            fresh = Archive.load(archive.dumps())
            incoming, more, hint = scan_path(Path(tmp))
            with self.assertRaises(Exception) as caught:
                fresh.apply_scan(incoming, more, None, hint)
            self.assertEqual(caught.exception.collisions[0]["label"], "tp1")
            self.assertEqual(len(fresh.blocks), 1)
            fresh.apply_scan(incoming, more, {"tp1": {"action": "overwrite"}}, hint)
            self.assertEqual(fresh.blocks[0].attributes["docker image"], "rocm/sgl-dev:test")
            incoming, more, hint = scan_path(Path(tmp))
            fresh.apply_scan(incoming, more, {"tp1": {"action": "rename", "label": "tp1 (2)"}}, hint)
            self.assertEqual([block.label for block in fresh.blocks], ["tp1", "tp1 (2)"])

    def test_shared_baseline_file_becomes_a_choice_on_each_block(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "base", 1024, 1024, 4)
            write_bench(tmp, "cand", 1024, 1024, 4)
            blocks, _notices, model = scan_path(Path(tmp))
            raw = json.loads(Archive(blocks=blocks, model=model).dumps())
            base_id = next(block["id"] for block in raw["blocks"] if block["label"] == "base")
            raw["baseline_id"] = base_id
            for block in raw["blocks"]:
                block.pop("baseline", None)
                block.pop("baseline_id", None)
            loaded = Archive.load(json.dumps(raw))
            by = {block.label: block for block in loaded.blocks}
            self.assertTrue(by["base"].baseline)
            self.assertIsNone(by["base"].baseline_id)
            self.assertEqual(by["cand"].baseline_id, base_id)
            self.assertEqual(loaded.present()["blocks"][loaded.blocks.index(by["cand"])]["baseline_label"], "base")
            raw["blocks"][1]["baseline"] = True
            raw["baseline_id"] = raw["blocks"][0]["id"]
            kept = Archive.load(json.dumps(raw))
            self.assertTrue(any(block.baseline for block in kept.blocks))
            self.assertFalse(any(block.baseline_id == raw["blocks"][0]["id"] for block in kept.blocks))

    def test_refs_stay_at_the_directory_and_retire_frees_the_label(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = root / "first"
            nested = first / "sweep"
            nested.mkdir(parents=True)
            write_bench(nested, "tp1", 1024, 1024, 4)
            second = root / "second"
            second.mkdir()
            write_bench(second, "other", 1024, 1024, 8)
            blocks, notices, model = scan_path(first)
            archive = Archive(blocks=blocks, notices=notices, model=model)
            added = archive.blocks[0].added_at
            self.assertRegex(added, r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}")
            self.assertEqual(archive.blocks[0].refs, [str(normalize(first))])
            self.assertNotIn(".json", archive.blocks[0].refs[0])
            more, extra_notices, _hint = scan_path(second)
            archive.adopt(archive.blocks[0].id, more, extra_notices, "")
            self.assertEqual(archive.blocks[0].refs, [str(normalize(first)), str(normalize(second))])
            third = root / "third"
            third.mkdir()
            write_bench(third, "later", 8192, 1024, 4)
            later, later_notices, _hint = scan_path(third)
            archive.blocks.append(later[0])
            archive.merge(later[0].id, archive.blocks[0].id, "keep")
            self.assertEqual(archive.blocks[0].refs, [str(normalize(first)), str(normalize(second)), str(normalize(third))])
            old = {
                "id": "old",
                "label": "tp1",
                "ingested_from": str(first),
                "points": [archive.blocks[0].points[0].to_dict()],
            }
            old["points"][0]["source"] = str(nested / "tp1_isl1024_osl1024_c4_r1.json")
            lifted = Archive.load(json.dumps({"format": "bta", "version": 1, "blocks": [old]}))
            self.assertEqual(lifted.blocks[0].refs, [str(first)])
            self.assertEqual(lifted.blocks[0].added_at, "")
            archive.patch(archive.blocks[0].id, {"retired": True, "baseline": True, "chart": True})
            self.assertTrue(archive.blocks[0].retired)
            self.assertFalse(archive.blocks[0].baseline)
            self.assertTrue(archive.blocks[0].retired_at)
            again, again_notices, hint = scan_path(first)
            archive.apply_scan(again, again_notices, None, hint)
            self.assertEqual([block.label for block in archive.blocks], ["tp1", "tp1"])
            self.assertEqual(archive.blocks[0].added_at, added)
            self.assertTrue(archive.blocks[1].added_at)
            self.assertNotEqual(archive.blocks[0].id, archive.blocks[1].id)
            self.assertTrue(archive.blocks[0].retired)
            self.assertFalse(archive.blocks[1].retired)
            with self.assertRaises(ArchiveError):
                archive.patch(archive.blocks[0].id, {"retired": False})
            archive.blocks[1].chart = True
            self.assertEqual(archive.present()["figures"][0]["series"][0]["label"], "tp1")
            self.assertEqual(len(archive.present()["figures"][0]["series"]), 1)
            data = workbook_bytes(archive)
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(data)).active
            labels = [sheet.cell(row, 2).value for row in range(1, sheet.max_row + 1)]
            self.assertEqual(labels.count("tp1"), 1)
            stored = json.loads(archive.dumps())
            self.assertEqual(
                stored["blocks"][0]["refs"],
                [str(normalize(first)), str(normalize(second)), str(normalize(third))],
            )
            reloaded = Archive.load(archive.dumps())
            self.assertEqual([block.retired for block in reloaded.blocks], [True, False])
            self.assertEqual(reloaded.blocks[0].added_at, added)
            self.assertEqual([item["added_at"] for item in reloaded.present()["blocks"]], [added, archive.blocks[1].added_at])

    def test_sheet_formulas_and_chart(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "MI355", 1024, 1024, 4, total_input_tokens=50, total_output_tokens=50)
            write_bench(tmp, "MI355", 1024, 1024, 8, total_input_tokens=50, total_output_tokens=50)
            write_bench(tmp, "GB300", 1024, 1024, 4, total_input_tokens=150, total_output_tokens=50)
            write_bench(tmp, "GB300", 1024, 1024, 8, total_input_tokens=750, total_output_tokens=50)
            blocks, _notices, model = scan_path(Path(tmp))
            by = {block.label: block for block in blocks}
            archive = Archive(model=model, blocks=[by["MI355"], by["GB300"]])
            archive.blocks[0].baseline = True
            archive.blocks[0].chart = True
            archive.blocks[0].attributes["chip"] = "MI355"
            archive.blocks[0].attributes["=chip"] = "=not-a-formula"
            archive.blocks[1].baseline_id = archive.blocks[0].id
            archive.blocks[1].chart = True
            archive.blocks[1].attributes["docker image"] = "image"
            archive.blocks[1].attributes["host name"] = "box"
            archive.blocks[1].attributes["accuracy score"] = "0.96"
            archive.blocks[1].attributes["BKC"] = "bkc-1"
            data = workbook_bytes(archive)
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data))
            sheet = book.active
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertIsNone(sheet.freeze_panes)
            self.assertEqual(sheet["B2"].value, "MI355")
            self.assertIsNone(sheet["F2"].value)
            self.assertEqual(sheet["E7"].data_type, "f")
            self.assertEqual(sheet["E7"].value, "=W7")
            self.assertEqual(sheet["F7"].value, '=IF(H7="","",1000/H7)')
            self.assertEqual(sheet["G7"].value, "=AA7")
            self.assertEqual(sheet["H7"].value, "=AE7")
            self.assertEqual(sheet["B11"].value, "GB300")
            self.assertEqual(sheet["I11"].value, "GB300 v.s. MI355")
            self.assertEqual(sheet["I16"].value, "=E16/E7")
            self.assertEqual(sheet["K17"].value, "=H17/H8")
            self.assertEqual(sheet["I18"].value, "=GEOMEAN(I16:I17)")
            self.assertEqual(sheet["D18"].value, "geomean")
            self.assertEqual(sheet["I13"].value, "BKC")
            self.assertEqual(sheet["J13"].value, "bkc-1")
            self.assertEqual(sheet["F13"].value, 0.96)
            self.assertEqual(len(sheet._charts), 1)
            titles = [series.title.v for series in sheet._charts[0].series]
            self.assertEqual(titles, ["MI355", "GB300"])
            odd = Archive.load(archive.dumps())
            self.assertEqual(odd.blocks[0].attributes["=chip"], "=not-a-formula")

    def test_ratio_gap_is_a_blank_cell(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "base", 1024, 1024, 4)
            write_bench(tmp, "base", 1024, 1024, 8)
            write_bench(tmp, "cand", 1024, 1024, 4)
            write_bench(tmp, "cand", 1024, 1024, 8)
            write_bench(tmp, "cand", 1024, 1024, 16)
            write_bench(tmp, "cand", 2048, 1024, 4)
            blocks, _notices, _model = scan_path(Path(tmp))
            by = {block.label: block for block in blocks}
            archive = Archive(blocks=[by["base"], by["cand"]])
            archive.blocks[0].baseline = True
            archive.blocks[1].baseline_id = archive.blocks[0].id
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(workbook_bytes(archive))).active
            self.assertIsNone(sheet["I18"].value)
            self.assertNotEqual(sheet["I18"].data_type, "inlineStr")
            self.assertEqual(sheet["I19"].value, "=GEOMEAN(I16:I18)")
            self.assertIsNone(sheet["I21"].value)
            self.assertIsNone(sheet["I22"].value)
            view = archive.present()
            shape = view["blocks"][1]["shapes"][0]
            self.assertIsNone(shape["rows"][2]["ratios"])
            self.assertAlmostEqual(shape["geomean"]["tput"], 1)
            self.assertIsNone(view["blocks"][1]["shapes"][1]["geomean"])
            next(point for point in archive.blocks[0].points if point.concurrency == 4).median_ttft_ms = 0
            gap = archive.present()["blocks"][1]["shapes"][0]["rows"][0]["ratios"]
            self.assertIn("tput", gap)
            self.assertNotIn("ttft", gap)
            self.assertIn("tpot", gap)

    def test_bad_json_and_later_framework_do_not_drop_the_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            good = bench_row(input_len=32, output_len=32, max_concurrency=4)
            (Path(tmp) / "keep_isl32_osl32_c4_r1.json").write_text(json.dumps(good), encoding="utf-8")
            broken = bench_row(input_len=32, output_len=32, max_concurrency=8)
            broken["input_len"] = None
            (Path(tmp) / "keep_isl32_osl32_c8_r1.json").write_text(json.dumps(broken), encoding="utf-8")
            log = Path(tmp) / "late_isl8_osl8_c4_r1.log"
            log.write_text(
                "Total token throughput (tok/s): 100\nMedian TTFT (ms): 10\nMedian TPOT (ms): 20\n",
                encoding="utf-8",
            )
            write_bench(tmp, "late", 8, 8, 8)
            blocks, notices, _model = scan_path(Path(tmp))
            by = {block.label: block for block in blocks}
            self.assertEqual(sorted(point.concurrency for point in by["keep"].points), [4])
            self.assertTrue(any("not an integer" in item["message"] and item["kind"] == "skip" for item in notices))
            self.assertEqual(by["late"].attributes["framework"], "sglang")

    def test_append_keeps_the_other_tabs(self):
        from openpyxl import Workbook, load_workbook
        book = Workbook()
        notes = book.active
        notes.title = "Notes"
        notes["A1"] = "keep"
        notes["B1"] = "=1+1"
        raw = io.BytesIO()
        book.save(raw)
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "MI355", 1024, 1024, 4)
            blocks, _notices, model = scan_path(Path(tmp))
            archive = Archive(model=model, blocks=blocks)
        data, name = append_workbook(raw.getvalue(), archive, "bta-1023-0700")
        again = load_workbook(io.BytesIO(data))
        self.assertEqual(again["Notes"]["A1"].value, "keep")
        self.assertEqual(again["Notes"]["B1"].value, "=1+1")
        self.assertEqual(name, "bta-1023-0700")
        self.assertEqual(again[name]["E7"].value, "=W7")
        self.assertEqual(again.active.title, name)
        data, name = append_workbook(data, archive, "bta-1023-0700")
        final = load_workbook(io.BytesIO(data))
        self.assertEqual(name, "bta-1023-0700 2")
        self.assertEqual(final.sheetnames, ["Notes", "bta-1023-0700 2", "bta-1023-0700"])
        self.assertEqual(final["Notes"]["A1"].value, "keep")

    def test_protected_workbook_is_described_without_its_bytes(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_bench(tmp, "base", 1024, 1024, 4)
            write_bench(tmp, "base", 1024, 1024, 8)
            write_bench(tmp, "cand", 1024, 1024, 4)
            write_bench(tmp, "cand", 1024, 1024, 8)
            blocks, _notices, model = scan_path(Path(tmp))
            by = {block.label: block for block in blocks}
            by["base"].baseline = True
            by["base"].chart = True
            by["cand"].baseline_id = by["base"].id
            by["cand"].chart = True
            archive = Archive(model=model, blocks=[by["base"], by["cand"]])
        plan = describe(workbook_bytes(archive, sheet_title="bta-1023-0700"))
        flat = [cell for chunk in plan["chunks"] for row in chunk["formulas"] for cell in row]
        self.assertIn("=W7", flat)
        self.assertTrue(any(isinstance(cell, str) and cell.startswith("=GEOMEAN(") for cell in flat))
        self.assertTrue(any(item.split(":")[0].startswith("I") and item.split(":")[-1].startswith("K") for item in plan["merges"]))
        self.assertTrue(any(item["body"]["color"] == "#92D050" for item in plan["fills"]))
        widths = {item["address"]: item["width"] for item in plan["widths"]}
        self.assertAlmostEqual(widths["A:A"], round(3 * 54 / 8.43, 2))
        self.assertGreater(widths["B:B"], 80)
        self.assertLess(len(plan["borders"]), 80)
        self.assertEqual(plan["base_font"]["body"]["name"], "Arial")
        self.assertEqual(len(plan["charts"]), 1)
        chart = plan["charts"][0]
        island = [cell for row in chart["formulas"] for cell in row]
        self.assertIn("base", island)
        self.assertIn("cand", island)
        self.assertIn("=NA()", island)
        self.assertIn("ISL", chart["title"])
        self.assertEqual(place_request("bta-1023-0700"), {"position": 1})
        message = _http_message(200, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(range(256)) * 4)
        self.assertEqual(message, "")
        self.assertNotIn("\ufffd", message)

    def test_onedrive_link_names_the_file(self):
        url = (
            "https://amdcloud-my.sharepoint.com/:x:/r/personal/akao_amd_com/"
            "Documents/grok2_5_akao.xlsx?d=w886a660630a241a1970a8fabfc2e74ff&csf=1&web=1"
        )
        self.assertEqual(file_name(url), "grok2_5_akao.xlsx")
        share = encode_share_url(url)
        self.assertTrue(share.startswith("u!"))
        self.assertNotIn("=", share)
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "onedrive.url"
            kept = remember_url("  https://amdcloud-my.sharepoint.com/x/other.xlsx  ", dest)
            self.assertEqual(dest.read_text(encoding="utf-8").strip(), kept)
            self.assertTrue(kept.endswith("other.xlsx"))
            with self.assertRaises(OneDriveError):
                remember_url("not a link", dest)
            self.assertTrue(dest.read_text(encoding="utf-8").strip().endswith("other.xlsx"))

    def test_locked_workbook_is_edited_through_excel(self):
        calls = []

        def zip_add():
            calls.append("zip")
            raise WorkbookLocked("locked")

        def excel_add():
            calls.append("excel")
            return {"sheet": "bta"}

        self.assertEqual(_open_for_edit(b"PK\x03\x04rest", zip_add, excel_add)["sheet"], "bta")
        self.assertEqual(calls, ["zip", "excel"])
        calls.clear()
        self.assertEqual(_open_for_edit(b"\xd0\xcf\x11\xe0rest", zip_add, excel_add)["sheet"], "bta")
        self.assertEqual(calls, ["excel"])

    def test_real_directory_if_present(self):
        if not REAL.is_dir():
            self.skipTest("sample sweep is not on this machine")
        blocks, notices, model = scan_path(REAL)
        self.assertEqual([block.label for block in blocks], ["tp1"])
        points = [point for point in blocks[0].points if point.run == 1]
        self.assertEqual(sorted(point.concurrency for point in points), [4, 8, 16, 32, 64])
        self.assertTrue(all(point.input_len == 70000 and point.output_len == 300 for point in points))
        self.assertTrue(model)
        self.assertFalse(any("throughput check failed" in item["message"] for item in notices))

    def test_file_list_is_newest_first(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            older = root / "older-dir"
            newer = root / "newer.txt"
            middle = root / "middle.log"
            older.mkdir()
            newer.write_text("x", encoding="utf-8")
            middle.write_text("y", encoding="utf-8")
            os.utime(older, (1_000_000_000, 1_000_000_000))
            os.utime(middle, (1_500_000_000, 1_500_000_000))
            os.utime(newer, (1_700_000_000, 1_700_000_000))
            names = [item["name"] for item in list_dir(str(root))["entries"]]
            self.assertEqual(names, ["newer.txt", "middle.log", "older-dir"])

    def test_repeated_underscores_and_optional_run(self):
        if not R1_JSON.is_dir():
            self.skipTest("r1 json folder is not on this machine")
        blocks, notices, model = scan_path(R1_JSON)
        by = {block.label: block for block in blocks}
        self.assertEqual(sorted(by), ["R1__TP1", "R1__TP2"])
        self.assertEqual(len(by["R1__TP1"].points), 15)
        self.assertEqual(len(by["R1__TP2"].points), 15)
        self.assertTrue(all(point.run == 1 for block in blocks for point in block.points))
        self.assertEqual(model, "gpt-oss-120b-w-mxfp4-a-fp8")
        self.assertFalse(notices)


class ApiTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.workspace = Workspace(self.root / "workspace.bta")
        self.app = App(self.workspace)
        self.httpd = serve(self.app, "127.0.0.1", 0)
        self.port = self.httpd.server_address[1]
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        self.tmp.cleanup()

    def request(self, method, path, body=None):
        data = None if body is None else json.dumps(body).encode()
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", data=data, method=method)
        if data is not None:
            req.add_header("Content-Type", "application/json")
        try:
            with urllib.request.urlopen(req) as response:
                raw = response.read()
                return response.status, json.loads(raw.decode() or "null") if "json" in response.headers.get("Content-Type", "") else raw
        except urllib.error.HTTPError as exc:
            payload = exc.read()
            try:
                payload = json.loads(payload.decode())
            except json.JSONDecodeError:
                pass
            return exc.code, payload

    def test_case_a_and_b(self):
        status, listing = self.request("GET", "/api/fs?path=" + urllib.parse.quote(str(self.root)))
        self.assertEqual(status, 200)
        self.assertEqual(listing["path"], str(self.root.resolve()))
        results = self.root / "results"
        results.mkdir()
        write_bench(results, "tp1", 1024, 1024, 4)
        write_bench(results, "tp1", 8192, 1024, 4)
        status, view = self.request("POST", "/api/act", {"action": "ingest", "path": str(results)})
        self.assertEqual(status, 200, view)
        self.assertEqual(view["blocks"][0]["label"], "tp1")
        self.assertEqual(len(view["blocks"][0]["shapes"]), 2)
        block_id = view["blocks"][0]["id"]
        extra = self.root / "extra"
        extra.mkdir()
        write_bench(extra, "elsewhere", 8192, 1024, 8)
        write_bench(extra, "elsewhere", 1024, 1024, 4)
        status, view = self.request("POST", "/api/act", {"action": "ingest", "path": str(extra), "into": block_id})
        self.assertEqual(status, 200, view)
        self.assertEqual([item["label"] for item in view["blocks"]], ["tp1"])
        shapes = {(shape["input_len"], shape["output_len"]): [row["concurrency"] for row in shape["rows"]] for shape in view["blocks"][0]["shapes"]}
        self.assertEqual(shapes[(1024, 1024)], [4])
        self.assertEqual(shapes[(8192, 1024)], [4, 8])
        self.assertTrue(any("already in tp1 1024/1024 c4 r1" in item["message"] for item in view["notices"]))
        solo = self.root / "solo"
        solo.mkdir()
        write_bench(solo, "solo", 70000, 300, 16)
        status, view = self.request("POST", "/api/act", {"action": "ingest", "path": str(solo)})
        self.assertEqual(status, 200, view)
        solo_id = next(item["id"] for item in view["blocks"] if item["label"] == "solo")
        status, view = self.request("POST", "/api/act", {"action": "merge", "id": solo_id, "into": block_id, "choice": "keep"})
        self.assertEqual(status, 200, view)
        self.assertEqual([item["label"] for item in view["blocks"]], ["tp1"])
        self.assertTrue(any(shape["input_len"] == 70000 for shape in view["blocks"][0]["shapes"]))
        status, view = self.request("POST", "/api/act", {
            "action": "patch",
            "id": block_id,
            "baseline": True,
            "chart": True,
            "gpus": 2,
            "attributes": {"docker image": "rocm/sgl", "host name": "m12", "chip": "MI355", "framework": "SGLang", "tp": "1", "accuracy score": "", "BKC": "", "extra": "kept"},
        })
        self.assertEqual(status, 200, view)
        block = view["blocks"][0]
        self.assertEqual(block["gpus"], 2)
        self.assertAlmostEqual(block["shapes"][0]["rows"][0]["per_gpu"], 500)
        status, denied = self.request("POST", "/api/act", {"action": "ingest", "path": str(results)})
        self.assertEqual(status, 409)
        self.assertEqual(denied["collisions"][0]["label"], "tp1")
        status, view = self.request("GET", "/api/workspace")
        self.assertEqual(len(view["blocks"]), 1)
        self.assertEqual(view["blocks"][0]["attributes"][0]["value"], "rocm/sgl")
        status, view = self.request("POST", "/api/act", {
            "action": "ingest",
            "path": str(results),
            "resolutions": {"tp1": {"action": "rename", "label": "tp1 copy"}},
        })
        self.assertEqual([block["label"] for block in view["blocks"]], ["tp1", "tp1 copy"])
        self.assertEqual(view["blocks"][0]["attributes"][0]["key"], "docker image")
        self.assertEqual(view["blocks"][0]["attributes"][0]["value"], "rocm/sgl")
        copy_id = view["blocks"][1]["id"]
        status, view = self.request("POST", "/api/act", {"action": "patch", "id": block_id, "label": "MI355 TP1"})
        status, view = self.request("POST", "/api/act", {"action": "patch", "id": copy_id, "baseline_id": block_id, "chart": True})
        self.assertEqual(view["blocks"][1]["baseline_label"], "MI355 TP1")
        self.assertEqual(len(view["figures"]), 3)
        status, xlsx = self.request("GET", "/api/export.xlsx")
        self.assertEqual(status, 200)
        self.assertTrue(xlsx[:2] == b"PK")
        status, text = self.request("GET", "/api/export.bta")
        saved = self.root / "saved.bta"
        saved.write_bytes(text)
        source = next(results.glob("*.json"))
        source.unlink()
        status, view = self.request("POST", "/api/act", {"action": "import", "path": str(saved)})
        self.assertEqual(status, 200, view)
        self.assertGreater(view["blocks"][0]["missing"], 0)
        self.assertEqual(view["blocks"][0]["label"], "MI355 TP1")
        status, view = self.request("POST", "/api/act", {"action": "delete", "id": copy_id})
        self.assertEqual(len(view["blocks"]), 1)
        page = urllib.request.urlopen(f"http://127.0.0.1:{self.port}/").read().decode()
        self.assertIn("benchmark-tracker", page)


if __name__ == "__main__":
    unittest.main()
